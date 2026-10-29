import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_user_guide():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Palette
    COLOR_PRIMARY = RGBColor(157, 23, 77)     # Crimson / Wine
    COLOR_GOLD = RGBColor(197, 155, 39)       # Gold
    COLOR_DARK = RGBColor(30, 41, 59)         # Slate Dark
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Muted
    COLOR_HEADER_BG = "9D174D"                # Wine Hex
    COLOR_LIGHT_BG = "F8FAFC"                 # Light Slate Hex
    COLOR_CALLOUT_BG = "FFFBEB"               # Warm Gold Hex
    COLOR_CALLOUT_BORDER = "F59E0B"           # Amber Hex

    # Document Header / Logo Section
    logo_path = r"c:\Users\Guest User\.gemini\antigravity-ide\scratch\bible_projector\alappuzha_cog_logo.png"
    if os.path.exists(logo_path):
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_logo = logo_para.add_run()
        run_logo.add_picture(logo_path, width=Inches(1.3))

    # Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_church = p_title.add_run("CHURCH OF GOD (FULL GOSPEL) IN INDIA\nALAPPUZHA TOWN • KERALA STATE")
    run_church.font.name = "Calibri"
    run_church.font.size = Pt(13)
    run_church.font.bold = True
    run_church.font.color.rgb = COLOR_PRIMARY

    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_main.paragraph_format.space_before = Pt(6)
    p_main.paragraph_format.space_after = Pt(4)
    run_main = p_main.add_run("ChurchView Pro Dual Bible Projector\nComplete Operator & Setup Guide")
    run_main.font.name = "Calibri"
    run_main.font.size = Pt(22)
    run_main.font.bold = True
    run_main.font.color.rgb = COLOR_DARK

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Malayalam BSI + English NKJV / NIV / KJV • 100% Offline Church Scripture System")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_GOLD

    # Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("—" * 58)
    run_div.font.color.rgb = COLOR_GOLD
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Helper function for Section Headings
    def add_section_header(title, num=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run_num = p.add_run(f"{num} " if num else "")
        run_num.font.name = "Calibri"
        run_num.font.size = Pt(14)
        run_num.font.bold = True
        run_num.font.color.rgb = COLOR_PRIMARY

        run_txt = p.add_run(title)
        run_txt.font.name = "Calibri"
        run_txt.font.size = Pt(14)
        run_txt.font.bold = True
        run_txt.font.color.rgb = COLOR_DARK
        return p

    def add_callout(text, title="IMPORTANT NOTE:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, "FFFBEB")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        # Border
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="F59E0B"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        r_title = cp.add_run(f"{title} ")
        r_title.bold = True
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = RGBColor(180, 83, 9)

        r_text = cp.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = COLOR_DARK
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Overview
    add_section_header("Overview & Core Capabilities", "1.")
    p_ov = doc.add_paragraph()
    p_ov.paragraph_format.space_after = Pt(8)
    p_ov.add_run(
        "ChurchView Pro is a professional, bilingual church scripture projection software tailored specifically for "
        "Church of God (Full Gospel) in India, Alappuzha Town. It allows church operators to project scripture verses "
        "simultaneously in Malayalam (BSI) and English (NKJV, NIV, or KJV) with zero latency onto the sanctuary projector or TV."
    )

    # Key Features Bullet points
    features = [
        ("100% Offline Operation", "Does not require an internet connection, Wi-Fi router, or mobile data during church service. All 31,102 verses and graphics are local."),
        ("Bilingual Dual Scripture", "Presents authentic Malayalam BSI along with your choice of English translation (NKJV, NIV, or KJV)."),
        ("Live Browsing Isolation", "Operators can browse different books and chapters freely on the laptop without changing what the congregation sees on the projector."),
        ("Arrow Key Smooth Navigation", "Press Right (→) or Down (↓) to advance verses; press Left (←) or Up (↑) to go back. Smoothly crosses chapters and books."),
        ("Authentic Church Holding Screen", "1-click toggle ('W' key) displays the beautiful Church of God Kerala State logo and Alappuzha Town welcome graphic."),
        ("Instant Scripture Search", "Jump straight to any verse in seconds by typing 'John 3:16', 'യോഹ 3:16', 'Rom 8:28', or 'Gen 1:1' and pressing Enter."),
        ("Recent Verses History Shelf", "Every projected verse is automatically saved in a 1-click recall shelf so you can instantly re-project verses referenced earlier in the sermon."),
        ("Scripture Watermark & Themes", "Configurable church watermark emblem and high-contrast background themes (Midnight Navy, Pure Black, Royal Burgundy, Deep Emerald, Warm Charcoal).")
    ]
    for feat_title, feat_desc in features:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(f"{feat_title}: ")
        r1.bold = True
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = bp.add_run(feat_desc)
        r2.font.color.rgb = COLOR_DARK

    add_callout(
        "ChurchView runs entirely from a local folder on your laptop without installing heavy external software, databases, or subscriptions. It runs on any Windows 10 or Windows 11 PC out-of-the-box.",
        "ZERO-DEPENDENCY PORTABILITY:"
    )

    # 2. Step-by-Step Installation
    add_section_header("How to Copy & Install on the Church Laptop", "2.")
    steps = [
        ("Step 2.1: Transfer the ZIP File to the Laptop", 
         "Copy the file 'ChurchView_Alappuzha_Town.zip' from your USB flash drive or Downloads folder onto the church laptop. We recommend placing it in your Documents, Desktop, or directly in C:\\ChurchView."),
        ("Step 2.2: Extract All Files", 
         "Right-click on 'ChurchView_Alappuzha_Town.zip' and select 'Extract All...'. Choose a destination folder and click 'Extract'. Ensure that all files (especially bible_dual.json and Launch_ChurchView_Projector.bat) remain together in that single unzipped folder."),
        ("Step 2.3: Verify the Unzipped Contents", 
         "Confirm that the unzipped folder contains:\n"
         "  • Launch_ChurchView_Projector.bat (1-Click Launcher)\n"
         "  • server.py & server_ps.ps1 (Local servers for Python or native Windows PowerShell)\n"
         "  • bible_dual.json (Full Malayalam BSI + NKJV + NIV + KJV database)\n"
         "  • index.html & projector.html (Presenter and Projector interfaces)\n"
         "  • presenter.js, projector.js & styles.css\n"
         "  • Church logos and emblems (alappuzha_cog_logo.png, alappuzha_emblem_trans.png)\n"
         "  • README_SETUP.txt")
    ]
    for s_title, s_desc in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r_st = p.add_run(s_title)
        r_st.bold = True
        r_st.font.size = Pt(11)
        r_st.font.color.rgb = COLOR_PRIMARY

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(6)
        r_sd = p2.add_run(s_desc)
        r_sd.font.color.rgb = COLOR_DARK

    # 3. How to Launch
    add_section_header("1-Click Launching Procedure", "3.")
    p_launch = doc.add_paragraph()
    p_launch.paragraph_format.space_after = Pt(6)
    p_launch.add_run(
        "Launching ChurchView requires just a single double-click:\n"
        "1. Open your unzipped folder and double-click "
    )
    r_bat = p_launch.add_run("Launch_ChurchView_Projector.bat")
    r_bat.bold = True
    r_bat.font.color.rgb = COLOR_PRIMARY
    p_launch.add_run(".\n2. A black terminal window will open and automatically start the local web server. It will detect whether Python is installed; if not, it seamlessly starts using the native Windows PowerShell engine.\n"
                    "3. Your default web browser (Edge, Chrome, Brave, etc.) will automatically open to ")
    r_url = p_launch.add_run("http://localhost:8765")
    r_url.bold = True
    p_launch.add_run(" displaying the Presenter Workbench.\n"
                    "4. ")
    r_keep = p_launch.add_run("Crucial:")
    r_keep.bold = True
    p_launch.add_run(" Keep the small black command window open or minimized during church service, as it powers the local app.")

    # 4. Connecting to the Sanctuary Projector / TV
    add_section_header("Connecting & Configuring the Projector (Dual Display)", "4.")
    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_after = Pt(6)
    p_proj.add_run(
        "To present cleanly in church without showing your operator controls to the congregation, follow these standard steps:\n\n"
        "1. Connect your laptop to the projector or TV using an HDMI or VGA cable.\n"
        "2. Press "
    )
    r_winp = p_proj.add_run("Windows Key + P")
    r_winp.bold = True
    r_winp.font.color.rgb = COLOR_PRIMARY
    p_proj.add_run(" on your keyboard.\n"
                   "3. Select ")
    r_ext = p_proj.add_run("Extend")
    r_ext.bold = True
    p_proj.add_run(" (do NOT select 'Duplicate'). This gives you two independent screens: your laptop screen for controlling verses and the projector for the congregation.\n"
                   "4. In the ChurchView Presenter window on your laptop, click the gold button: ")
    r_btn = p_proj.add_run("OPEN PROJECTOR (O)")
    r_btn.bold = True
    r_btn.font.color.rgb = COLOR_GOLD
    p_proj.add_run(" in the top navigation bar.\n"
                   "5. A clean, borderless audience projector window will open.\n"
                   "6. Drag this projector window across to the second screen (the projector/TV).\n"
                   "7. Press ")
    r_f11 = p_proj.add_run("F11")
    r_f11.bold = True
    r_f11.font.color.rgb = COLOR_PRIMARY
    p_proj.add_run(" (or double-click anywhere on the screen) to make it Full Screen! The mouse cursor will automatically hide on the projector screen.")

    # 5. Live Operation: Browsing vs Projecting
    add_section_header("Operator Workflow: Browsing vs. Projecting Verses", "5.")
    p_op = doc.add_paragraph()
    p_op.paragraph_format.space_after = Pt(6)
    p_op.add_run(
        "One of the greatest features in ChurchView is the separation between browsing and live projection:\n\n"
        "• Browsing Books & Chapters: When the pastor mentions another scripture passage ahead of time, "
        "you can click any book in the left sidebar or select any chapter from the dropdown / pills. "
        "The verses will load immediately in the center panel for you to read and prepare, "
    )
    r_alert = p_op.add_run("BUT the congregation's screen will NOT change.")
    r_alert.bold = True
    r_alert.font.color.rgb = COLOR_PRIMARY
    p_op.add_run(" It remains completely locked on the previous live verse until you are ready.\n\n"
                   "• Projecting a Verse: To display a verse to the congregation, simply ")
    r_clk = p_op.add_run("click directly on that verse card")
    r_clk.bold = True
    p_op.add_run(". It will instantly highlight with a gold border and 'LIVE' badge on your workbench, and appear immediately on the church projector.\n\n"
                   "• Visual Status Badge: The header bar at the top of the verse list always displays the exact live status:\n"
                   "   - If viewing the active chapter: ")
    r_lh = p_op.add_run("🔴 LIVE: v.16")
    r_lh.bold = True
    r_lh.font.color.rgb = COLOR_PRIMARY
    p_op.add_run("\n   - If browsing another chapter: ")
    r_lo = p_op.add_run("📺 Screen: John 3:16")
    r_lo.bold = True
    r_lo.font.color.rgb = RGBColor(59, 130, 246)
    p_op.add_run(" (allowing you to know what is live without glancing away).")

    # 6. Keyboard Shortcuts Reference Table
    add_section_header("Keyboard Shortcuts Quick Reference Card", "6.")
    
    shortcuts = [
        ("→ (Right Arrow)", "Next Verse", "Advances sequentially to the next verse (auto-advances across chapters and books)"),
        ("↓ (Down Arrow)", "Next Verse", "Alternative key to advance to the next verse"),
        ("Spacebar", "Next Verse", "Alternative key to advance to the next verse"),
        ("← (Left Arrow)", "Previous Verse", "Moves back sequentially to the previous verse"),
        ("↑ (Up Arrow)", "Previous Verse", "Alternative key to move back to the previous verse"),
        ("W", "Default / Holding Screen", "Toggles the Church of God Alappuzha Town welcome holding screen"),
        ("B", "Blackout / Blank", "Blanks the audience projector screen completely (blackout mode)"),
        ("O", "Open Projector Window", "Opens or focuses the clean audience projector screen for the TV"),
        ("/ (Slash)", "Quick Search Box", "Directly focuses the search box to type scripture references (e.g. 'John 3:16')"),
        ("F11 or F", "Full Screen", "Toggles full screen mode on the audience projector window")
    ]

    table = doc.add_table(rows=len(shortcuts) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    headers = ["Key / Shortcut", "Action", "Description"]
    col_widths = [Inches(1.8), Inches(1.8), Inches(3.2)]
    
    hdr_row = table.rows[0]
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, COLOR_HEADER_BG)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, (key, act, desc) in enumerate(shortcuts, start=1):
        row = table.rows[r_idx]
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([key, act, desc]):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            elif c_idx == 1:
                r.bold = True
                r.font.color.rgb = COLOR_DARK
            else:
                r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 7. Translation & Layout Controls
    add_section_header("Scripture Translations & Layout Options", "7.")
    p_trans = doc.add_paragraph()
    p_trans.paragraph_format.space_after = Pt(6)
    p_trans.add_run(
        "ChurchView gives the operator total flexibility over scripture display:\n\n"
        "1. English Version Selection: Use the dropdown in the top right header to switch between "
    )
    r_nkjv = p_trans.add_run("NKJV (New King James Version)")
    r_nkjv.bold = True
    p_trans.add_run(", ")
    r_niv = p_trans.add_run("NIV (New International Version)")
    r_niv.bold = True
    p_trans.add_run(", or ")
    r_kjv = p_trans.add_run("KJV (King James Version)")
    r_kjv.bold = True
    p_trans.add_run(". The change reflects instantly across both workbench and projector.\n\n"
                   "2. Scripture Layouts (Under Projector Display Settings in right sidebar):\n"
                   "   • Dual (Malayalam + English): Standard church mode displaying Malayalam BSI on top with the English translation underneath.\n"
                   "   • Malayalam BSI Only: Projects only the Malayalam scripture in large, prominent font size (ideal for Malayalam-only services).\n"
                   "   • English Only: Projects only the English translation in large font size (ideal for English or youth services).\n\n"
                   "3. Font Size Sliders: Real-time sliders let you fine-tune the Malayalam and English font sizes to match the exact size and aspect ratio of your church screen.")

    # 8. Customizing the Default Holding Screen
    add_section_header("Customizing the Default / Holding Screen", "8.")
    p_def = doc.add_paragraph()
    p_def.paragraph_format.space_after = Pt(6)
    p_def.add_run(
        "When welcoming the congregation before service starts or during prayer transitions, pressing 'W' displays the authentic Church of God holding screen.\n\n"
        "You can customize this screen anytime:\n"
        "1. In the right panel under 'Projector Display Settings', click "
    )
    r_ed = p_def.add_run("✏️ Edit Default Screen Text & Logo")
    r_ed.bold = True
    r_ed.font.color.rgb = COLOR_PRIMARY
    p_def.add_run(".\n2. An interactive customization popup will appear, allowing you to edit:\n"
                   "   • Organization Title (Default: CHURCH OF GOD (FULL GOSPEL) IN INDIA)\n"
                   "   • Branch / Subtitle (Default: ALAPPUZHA TOWN)\n"
                   "   • Service Motto (e.g. WELCOME • WORSHIP • FELLOWSHIP)\n"
                   "   • Featured Scripture Quote (e.g. Psalm 150:6)\n"
                   "   • Church Logo Image (upload your own image or emblem anytime)\n"
                   "3. Click 'Save & Apply'. The updated design will immediately save to your browser's local memory and reflect on the screen whenever 'W' is activated.")

    # 9. Recent History Bar
    add_section_header("Recent Verses Shelf (1-Click Recall)", "9.")
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_after = Pt(6)
    p_rec.add_run(
        "During expository sermons, pastors frequently cross-reference scriptures or return to earlier verses. "
        "ChurchView automatically records every verse you project into the '🕒 RECENT' bar located directly below the chapter selector.\n\n"
        "• To return to any previously projected verse, simply click its chip in the recent bar (e.g., [John 3:16], [Rom 8:28]).\n"
        "• The system instantly jumps back to that chapter and projects the verse in one click.\n"
        "• To reset the shelf for a new service, click the small '✕' button next to the history bar.")

    # 10. Troubleshooting & FAQ
    add_section_header("Troubleshooting & Operator FAQ", "10.")
    faqs = [
        ("Q1: Does this require internet at the church?",
         "A: Absolutely not. ChurchView is 100% offline. All 31,102 verses, Malayalam BSI, NKJV, NIV, KJV, graphics, and server code run completely on your local laptop."),
        ("Q2: What if the black launcher window closes?",
         "A: If the command window closes, the local app will stop. Simply double-click 'Launch_ChurchView_Projector.bat' again to restart it. It takes only 2 seconds."),
        ("Q3: What if the audience projector screen closes accidentally?",
         "A: Simply click 'OPEN PROJECTOR (O)' in the top header or press the 'O' key on your keyboard. It will immediately re-open and re-sync with whatever verse is currently live."),
        ("Q4: Why does my laptop show the same thing as the projector?",
         "A: You have Windows display set to 'Duplicate'. Press 'Windows Key + P' and choose 'Extend'. Then drag the projector window over to the projector screen and press F11."),
        ("Q5: What if port 8765 is reported in use?",
         "A: A previous instance of the server may still be running in the background. The launcher will automatically detect this and open the browser. Alternatively, close any existing black terminal windows and relaunch."),
        ("Q6: How do I transfer updates or custom logos to another PC?",
         "A: Simply copy the unzipped folder (or re-zip it) onto your USB drive and take it to the other machine. Everything is contained within that single folder.")
    ]
    for q, a in faqs:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(q)
        r_q.bold = True
        r_q.font.color.rgb = COLOR_PRIMARY

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_after = Pt(6)
        r_a = p_a.add_run(a)
        r_a.font.color.rgb = COLOR_DARK

    # Footer note
    p_ft = doc.add_paragraph()
    p_ft.paragraph_format.space_before = Pt(16)
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ft = p_ft.add_run("Church of God (Full Gospel) in India, Alappuzha Town • Media & Projection Ministry")
    r_ft.font.name = "Calibri"
    r_ft.font.size = Pt(10)
    r_ft.font.bold = True
    r_ft.font.color.rgb = COLOR_MUTED

    # Save to local path and Downloads
    out_local = r"c:\Users\Guest User\.gemini\antigravity-ide\scratch\bible_projector\ChurchView_User_Guide_Alappuzha_Town.docx"
    out_downloads = r"C:\Users\Guest User\Downloads\ChurchView_User_Guide_Alappuzha_Town.docx"

    doc.save(out_local)
    doc.save(out_downloads)
    print(f"Successfully generated: {out_local} and {out_downloads}")

if __name__ == '__main__':
    create_user_guide()
